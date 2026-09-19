#!/usr/bin/env python3
"""Render the academic data-methodology Markdown as a polished Word report."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "KONBAUNG_DATA_AND_EMBEDDINGS_METHODOLOGY.md"
OUTPUT = ROOT / "KONBAUNG_DATA_AND_EMBEDDINGS_METHODOLOGY.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202A35"
MUTED = "667381"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F4F6F8"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    """Apply a solid background fill to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    """Set compact, readable table-cell margins in twentieths of a point."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    """Mark the first table row to repeat when a table crosses pages."""
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    """Keep a heading or callout label with the paragraph that follows it."""
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if value and node is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and node is not None:
        p_pr.remove(node)


def set_keep_together(paragraph, value: bool = True) -> None:
    """Prevent short paragraphs from splitting across a page when possible."""
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepLines"))
    if value and node is None:
        p_pr.append(OxmlElement("w:keepLines"))
    elif not value and node is not None:
        p_pr.remove(node)


def set_widow_control(paragraph) -> None:
    """Enable Word's widow/orphan control for body text."""
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:widowControl")) is None:
        p_pr.append(OxmlElement("w:widowControl"))


def add_page_number(paragraph) -> None:
    """Insert a live PAGE field in a footer paragraph."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)


def set_font(run, name: str = "Calibri", size: float | None = None, color: str | None = None) -> None:
    """Apply Latin and Myanmar-capable font declarations to a run."""
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), "Myanmar Text")
    r_fonts.set(qn("w:cs"), "Myanmar Text")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str, *, base_size: float = 11, color: str = INK) -> None:
    """Render Markdown emphasis, inline code, and literal text into one paragraph."""
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")
    cursor = 0
    for match in token_re.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_font(run, size=base_size, color=color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Consolas", size=max(8.5, base_size - 1), color=DARK_BLUE)
            run._element.get_or_add_rPr().append(OxmlElement("w:noProof"))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=base_size, color=color)
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=base_size, color=color)
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run, size=base_size, color=color)


def configure_styles(document: Document) -> None:
    """Install the compact-reference typography used throughout the report."""
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


def configure_section(section) -> None:
    """Apply US Letter geometry and restrained running furniture to a section."""
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("KONBAUNG DATASET  ·  METHODOLOGY & REPRODUCIBILITY")
    set_font(run, size=8.5, color=MUTED)
    run.bold = True

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("DIGHUM research-data report  ·  31 August 2026  ·  ")
    set_font(run, size=8.5, color=MUTED)
    add_page_number(paragraph)
    for footer_run in paragraph.runs:
        set_font(footer_run, size=8.5, color=MUTED)


def add_cover(document: Document) -> None:
    """Create a restrained editorial cover using Word-native elements only."""
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(24)
    run = paragraph.add_run("METHODOLOGY REPORT")
    set_font(run, size=10, color=BLUE)
    run.bold = True

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(74)
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run("Konbaung Historical-Relations\nDataset and Embedding Tables")
    set_font(run, size=28, color=INK)
    run.bold = True

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(22)
    run = paragraph.add_run("Construction, normalization, similarity analysis, and reproducibility")
    set_font(run, size=14, color=DARK_BLUE)

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_BLUE)
    set_cell_margins(cell, top=180, start=220, bottom=180, end=220)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(
        "Research-data methods for the DIGHUM statistical-analysis archive. "
        "The report focuses on the database, embeddings, similarity tables, and "
        "exploratory clusters; application software is intentionally excluded."
    )
    set_font(run, size=11, color=INK)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(150)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run("VERSION 1.0")
    set_font(run, size=9, color=MUTED)
    run.bold = True
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run("31 August 2026")
    set_font(run, size=11, color=INK)
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Prepared as an article-facing methodology and reproducibility record")
    set_font(run, size=9.5, color=MUTED)

    document.add_page_break()


def add_contents(document: Document) -> None:
    """Add a compact contents map that remains stable without dynamic fields."""
    heading = document.add_paragraph(style="Heading 1")
    heading.add_run("Contents")
    sections = [
        "Scope and study design",
        "Source text, segmentation, and translation",
        "Open-coded triples and cross-page deduplication",
        "Axial analytical categories",
        "Eight-view embeddings and type-level feature tables",
        "Similarity tables and exploratory clusters",
        "Archive construction and validation",
        "Statistical guidance, limitations, and reproducibility",
        "Appendices: analytical files and parameters",
    ]
    for item in sections:
        paragraph = document.add_paragraph(style="List Bullet")
        add_inline(paragraph, item, base_size=10.5)
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(16)
    note.paragraph_format.left_indent = Inches(0.25)
    note.paragraph_format.right_indent = Inches(0.25)
    add_inline(
        note,
        "Interpretive rule: the axial categories are supervised analytical codes; "
        "the embedding-derived communities are separate exploratory outputs.",
        base_size=10.5,
        color=DARK_BLUE,
    )
    set_keep_together(note)
    document.add_page_break()


def add_markdown_table(document: Document, rows: list[list[str]]) -> None:
    """Render a Markdown table with explicit width, margins, and repeating header."""
    if not rows:
        return
    width = len(rows[0])
    table = document.add_table(rows=1, cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    usable = 6.5
    if width == 3:
        proportions = [0.30, 0.17, 0.53]
    elif width == 4:
        proportions = [0.23, 0.23, 0.20, 0.34]
    else:
        proportions = [1 / width] * width
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.width = Inches(usable * proportions[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, PALE_BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        add_inline(paragraph, value, base_size=9.25, color=DARK_BLUE)
        for run in paragraph.runs:
            run.bold = True
    set_repeat_table_header(table.rows[0])

    for row_values in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.width = Inches(usable * proportions[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, value, base_size=9.1)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    """Read one well-formed Markdown pipe table and return its rows and next index."""
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
            rows.append(values)
        index += 1
    return rows, index


def render_markdown(document: Document, markdown: str) -> None:
    """Render the report body using native headings, lists, tables, and code blocks."""
    lines = markdown.splitlines()
    index = 0
    skip_front = True
    page_break_before = {
        "6. Eight-view embedding generation",
        "10. Archive construction",
        "11. Guidance for statistical analysis",
        "13. Reproducibility sequence and provenance files",
        "Appendix A. Primary analytical files",
    }

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if skip_front:
            if stripped == "### Scope":
                skip_front = False
            else:
                index += 1
                continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = parse_table(lines, index)
            add_markdown_table(document, rows)
            continue

        if stripped.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            index += 1
            table = document.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.autofit = False
            table.columns[0].width = Inches(6.25)
            cell = table.cell(0, 0)
            set_cell_shading(cell, PALE_GRAY)
            set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run("\n".join(code_lines))
            set_font(run, name="Consolas", size=9, color=INK)
            run._element.get_or_add_rPr().append(OxmlElement("w:noProof"))
            document.add_paragraph().paragraph_format.space_after = Pt(1)
            continue

        if stripped.startswith("## ") or stripped.startswith("### "):
            level = 1 if stripped.startswith("## ") else 2
            text = stripped[3:] if level == 1 else stripped[4:]
            if text == "Scope":
                level = 1
            if text in page_break_before:
                document.add_page_break()
            paragraph = document.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, text, base_size=16 if level == 1 else 13, color=BLUE)
            for run in paragraph.runs:
                run.bold = True
            set_keep_with_next(paragraph)
            index += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, text)
            set_widow_control(paragraph)
            index += 1
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            set_widow_control(paragraph)
            index += 1
            continue

        if stripped.startswith("\\["):
            equation_lines = [stripped]
            index += 1
            while index < len(lines) and not lines[index].strip().endswith("\\]"):
                equation_lines.append(lines[index].strip())
                index += 1
            if index < len(lines):
                equation_lines.append(lines[index].strip())
                index += 1
            equation = " ".join(equation_lines).replace("\\[", "").replace("\\]", "")
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(equation)
            set_font(run, name="Cambria Math", size=11, color=DARK_BLUE)
            run.italic = True
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate:
                index += 1
                break
            if (
                candidate.startswith(("## ", "### ", "- ", "```", "|", "\\["))
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))
        set_widow_control(paragraph)


def main() -> None:
    """Build the final DOCX from the audited Markdown source."""
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    document = Document()
    configure_styles(document)
    configure_section(document.sections[0])

    properties = document.core_properties
    properties.title = "Konbaung Historical-Relations Dataset and Embedding Tables"
    properties.subject = "Methodology and reproducibility report"
    properties.keywords = "Konbaung; historical relations; embeddings; axial coding; methodology"
    properties.comments = "Generated from the audited local methodology Markdown source."

    add_cover(document)
    add_contents(document)
    render_markdown(document, SOURCE.read_text(encoding="utf-8"))

    for paragraph in document.paragraphs:
        set_widow_control(paragraph)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

